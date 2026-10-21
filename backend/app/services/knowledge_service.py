"""Knowledge-base service (design doc 3.4): law/template management, vectorization into ES, hit stats."""
import contextlib
from datetime import date, datetime

from sqlalchemy.orm import Session

from app.ai.embedding_client import embed_texts
from app.es.client import get_es_client
from app.es.indices import INDEX_CONTRACT, INDEX_LEGAL, INDEX_TEMPLATE
from app.models.knowledge import ContractTemplate, LawRegulation


def _es():
    return get_es_client()


def _parse_date(s: str | None) -> date | None:
    if not s:
        return None
    try:
        return datetime.strptime(s, "%Y-%m-%d").date()
    except ValueError:
        return None


def index_law(db: Session, law: LawRegulation) -> bool:
    """Vectorize a law/case into legal_knowledge; return False on failure (vectorized stays False)."""
    es = _es()
    if es is None:
        return False
    try:
        vector = embed_texts([f"{law.title}\n{law.content}"])[0]
        doc_id = f"law_{law.id}"
        es.index(
            index=INDEX_LEGAL,
            id=doc_id,
            document={
                "doc_type": law.doc_type,
                "ref_id": law.id,
                "title": law.title,
                "content": law.content,
                "content_vector": vector,
                "effective_date": law.effective_date.isoformat() if law.effective_date else None,
                "expiry_date": law.expiry_date.isoformat() if law.expiry_date else None,
                "status": law.status,
                "source_ref": law.source_ref or "",
            },
        )
        law.es_doc_id = doc_id
        law.vectorized = True
        db.commit()
        return True
    except Exception:
        law.vectorized = False
        db.commit()
        return False


def delete_law_from_es(law: LawRegulation) -> None:
    es = _es()
    if es is None or not law.es_doc_id:
        return
    # elasticsearch-py 8+ removed delete(ignore=...); must use options(ignore_status=404)
    with contextlib.suppress(Exception):
        es.options(ignore_status=404).delete(index=INDEX_LEGAL, id=law.es_doc_id)


def index_template(db: Session, tpl: ContractTemplate) -> bool:
    es = _es()
    if es is None:
        return False
    try:
        vector = embed_texts([f"{tpl.name}（{tpl.contract_type}）\n{tpl.content}"])[0]
        doc_id = f"tpl_{tpl.id}_v{tpl.version}"
        es.index(
            index=INDEX_TEMPLATE,
            id=doc_id,
            document={
                "doc_type": "template",
                "ref_id": tpl.id,
                "title": tpl.name,
                "content": tpl.content,
                "content_vector": vector,
                "contract_type": tpl.contract_type,
                "clause_level": tpl.clause_level,
                "status": "active",
            },
        )
        tpl.es_doc_id = doc_id
        tpl.vectorized = True
        db.commit()
        return True
    except Exception:
        tpl.vectorized = False
        db.commit()
        return False


def delete_template_from_es(tpl: ContractTemplate) -> None:
    es = _es()
    if es is None or not tpl.es_doc_id:
        return
    with contextlib.suppress(Exception):
        es.options(ignore_status=404).delete(index=INDEX_TEMPLATE, id=tpl.es_doc_id)


def bulk_index_contract_clauses(docs: list[dict]) -> None:
    """Index contract clauses into contract_index (BM25 only, no vectors)."""
    es = _es()
    if es is None or not docs:
        return
    try:
        from elasticsearch import helpers

        actions = [
            {
                "_index": INDEX_CONTRACT,
                "_id": f"clause_{d['clause_id']}",
                "_source": d,
            }
            for d in docs
        ]
        helpers.bulk(es, actions)
    except Exception as e:
        from loguru import logger

        logger.warning(f"contract clause index failed: {e}")


def parse_law_rows(rows: list[dict]) -> list[LawRegulation]:
    """Excel row -> LawRegulation object (docType/title/content/sourceRef/effectiveDate/expiryDate/status)."""
    laws = []
    for r in rows:
        doc_type = str(r.get("docType") or r.get("doc_type") or "law").strip().lower()
        if doc_type not in ("law", "case"):
            doc_type = "law"
        status = str(r.get("status") or "active").strip().lower()
        if status not in ("active", "expired"):
            status = "active"
        laws.append(
            LawRegulation(
                doc_type=doc_type,
                title=str(r.get("title") or "").strip()[:255],
                content=str(r.get("content") or "").strip(),
                source_ref=str(r.get("sourceRef") or r.get("source_ref") or "").strip()[:255],
                effective_date=_parse_date(str(r.get("effectiveDate") or r.get("effective_date") or "") or None),
                expiry_date=_parse_date(str(r.get("expiryDate") or r.get("expiry_date") or "") or None),
                status=status,
                hit_count=0,
                vectorized=False,
                created_at=datetime.now(),
            )
        )
    return [law for law in laws if law.title and law.content]


def read_law_excel(filepath: str) -> list[dict]:
    """openpyxl parse: first row is header (docType/title/content/sourceRef/effectiveDate/expiryDate/status)."""
    from openpyxl import load_workbook

    wb = load_workbook(filepath, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = [str(c or "").strip() if c is not None else "" for c in next(rows_iter)]
    except StopIteration:
        return []
    # support Chinese headers
    zh_map = {"类型": "docType", "标题": "title", "内容": "content", "来源": "sourceRef",
              "生效日期": "effectiveDate", "失效日期": "expiryDate", "状态": "status"}
    header = [zh_map.get(h, h) for h in header]
    rows = []
    for values in rows_iter:
        if values is None or all(v is None or str(v).strip() == "" for v in values):
            continue
        row = {}
        for i, key in enumerate(header):
            if key and i < len(values):
                row[key] = values[i]
        rows.append(row)
    wb.close()
    return rows
