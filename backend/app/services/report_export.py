"""Review report export (design doc 3.2.1): Word (python-docx) / PDF (reportlab + system CJK font)."""
import io
from pathlib import Path

from docx import Document
from docx.shared import Pt
from loguru import logger

from app.models.review import ReviewReport, ReviewTask

LEVEL_LABELS = {"high": "高风险", "medium": "中风险", "low": "低风险", "none": "无风险"}


def _find_cjk_font() -> str | None:
    for name in ("msyh.ttc", "simsun.ttc", "simhei.ttf", "simkai.ttf"):
        p = Path("C:/Windows/Fonts") / name
        if p.exists():
            return str(p)
    return None


def export_word(report: ReviewReport, task: ReviewTask, findings) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10.5)

    doc.add_heading(f"AI 合同合规审查报告（任务 #{task.id}）", 0)
    doc.add_paragraph(f"合同：{task.contract.title if task.contract else task.contract_id}")
    doc.add_paragraph(
        f"审查模型：{task.model_version}　Prompt 版本：{task.prompt_version}　"
        f"耗时：{task.duration_ms or 0} ms　Token 消耗：{task.tokens_used or 0}"
    )
    doc.add_paragraph(
        f"风险统计：高风险 {report.high_count} / 中风险 {report.medium_count} / 低风险 {report.low_count}；"
        f"合规评分：{report.score} / 100"
    )
    doc.add_heading("总体结论", level=1)
    doc.add_paragraph(report.overall_opinion)

    doc.add_heading("风险发现清单", level=1)
    for i, f in enumerate(findings, start=1):
        doc.add_heading(
            f"{i}. [{LEVEL_LABELS.get(f.risk_level, f.risk_level)}] {f.clause_no} {f.risk_type}（来源：{'规则引擎' if f.source == 'rule' else '大模型 RAG'}）",
            level=2,
        )
        doc.add_paragraph(f"原文摘录：{f.original_text or '-'}")
        doc.add_paragraph(f"风险描述：{f.issue}")
        doc.add_paragraph(f"法律/模板依据：{f.legal_basis or '-'}")
        doc.add_paragraph(f"修改建议：{f.suggestion or '-'}")
        if f.suggested_text:
            doc.add_paragraph(f"建议条款全文：{f.suggested_text}")

    doc.add_paragraph("")
    doc.add_paragraph("本报告由 AI 智能合同合规审查平台自动生成，仅供内部合规参考。")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_pdf(report: ReviewReport, task: ReviewTask, findings) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    font_path = _find_cjk_font()
    font_name = "Helvetica"
    if font_path:
        try:
            pdfmetrics.registerFont(TTFont("CJK", font_path))
            font_name = "CJK"
        except Exception as e:
            logger.warning(f"register cjk font failed: {e}")

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title="AI 合同合规审查报告")
    h1 = ParagraphStyle("h1", fontName=font_name, fontSize=15, spaceAfter=8, spaceBefore=4)
    h2 = ParagraphStyle("h2", fontName=font_name, fontSize=12, spaceAfter=4, spaceBefore=8)
    body = ParagraphStyle("body", fontName=font_name, fontSize=10, leading=15)

    story = [
        Paragraph(f"AI 合同合规审查报告（任务 #{task.id}）", h1),
        Paragraph(f"合同：{task.contract.title if task.contract else task.contract_id}", body),
        Paragraph(
            f"审查模型：{task.model_version}　Prompt：{task.prompt_version}　耗时：{task.duration_ms or 0} ms", body
        ),
        Paragraph(
            f"风险统计：高 {report.high_count} / 中 {report.medium_count} / 低 {report.low_count}；评分 {report.score}/100", body
        ),
        Spacer(1, 6),
        Paragraph("总体结论", h1),
        Paragraph(report.overall_opinion, body),
        Spacer(1, 6),
        Paragraph("风险发现清单", h1),
    ]
    for i, f in enumerate(findings, start=1):
        story.append(
            Paragraph(
                f"{i}. [{LEVEL_LABELS.get(f.risk_level, f.risk_level)}] {f.clause_no} {f.risk_type}", h2
            )
        )
        story.append(Paragraph(f"原文：{f.original_text or '-'}", body))
        story.append(Paragraph(f"风险：{f.issue}", body))
        story.append(Paragraph(f"依据：{f.legal_basis or '-'}", body))
        story.append(Paragraph(f"建议：{f.suggestion or '-'}", body))
    story.append(Spacer(1, 10))
    story.append(Paragraph("本报告由 AI 智能合同合规审查平台自动生成，仅供内部合规参考。", body))
    doc.build(story)
    return buf.getvalue()
