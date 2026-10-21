r"""Offline audit reproductions. Uses only isolated SQLite, temporary storage and synthetic data.

Run: .venv\Scripts\python.exe audit_reproduce.py
No production database, LLM, ES or Celery access is needed.
"""
import json
import faulthandler
import socket
from pathlib import Path
from unittest.mock import patch

# Reuse SQLite infrastructure, but deliberately retain real JWT authentication.
from tests import conftest as infra
from tests import helpers
from fastapi.testclient import TestClient
from app.main import app
from app.db.session import get_db
from app.core.security import create_access_token
from app.models.user import User
from app.services import workflow_engine, parse_service, review_service
from app.models.contract import ContractClause
from docx import Document

results = []


def record(name, expected, actual):
    results.append({"name": name, "expected": expected, "actual": actual,
                    "issue_reproduced": actual != expected})


def override_db():
    session = infra.TestSessionLocal()
    try:
        yield session
    finally:
        session.close()


_connect = socket.socket.connect
_connect_ex = socket.socket.connect_ex


def no_network(sock, address):
    # Windows asyncio requires a loopback socketpair to implement its event loop.
    if isinstance(address, tuple) and address[0] in ("127.0.0.1", "::1"):
        return _connect(sock, address)
    raise RuntimeError("Network disabled for audit")


def no_network_ex(sock, address):
    if isinstance(address, tuple) and address[0] in ("127.0.0.1", "::1"):
        return _connect_ex(sock, address)
    raise RuntimeError("Network disabled for audit")


def main():
    infra._db_base.Base.metadata.drop_all(infra.test_engine)
    infra._db_base.Base.metadata.create_all(infra.test_engine)
    db = infra.TestSessionLocal()
    # workflow:start + review:resolve: the eligibility/amendment checks must reach their logic,
    # while "unprivileged workflow/prompt edit" still rely on missing workflow:define / system:prompt.
    role = helpers.make_role(db, code="ordinary", data_scope="self",
                             permissions=[helpers.make_permission(db, code="workflow:start"),
                                          helpers.make_permission(db, code="review:resolve")])
    alice = helpers.make_user(db, username="alice", roles=[role])
    bob = helpers.make_user(db, username="bob")
    contract = helpers.make_contract(db, owner_id=bob.id, status="review_passed")
    version = helpers.make_version(db, contract.id, parse_status="parsed")
    clause = helpers.make_clause(db, version.id, contract.id, content="private phone 13812345678")
    task = helpers.make_review_task(db, contract.id, version_id=version.id, status="completed")
    report = helpers.make_report(db, task.id, contract.id, score=100,
                                report_json={"summary": {"failedClauses": 1}})
    finding = helpers.make_finding(db, task.id, clause_id=clause.id, suggested_text="changed clause")
    prompt = helpers.make_prompt(db)
    helpers.make_trace(db, task_type="qa", ref_id=0, payload={"question": "bob private question"})
    definition = helpers.make_definition(db, graph=helpers.approval_graph())
    db.commit()
    token = create_access_token(alice.id, alice.username)
    app.dependency_overrides[get_db] = override_db
    client = TestClient(app, headers={"Authorization": "Bearer " + token})
    record("unauthenticated control", 401,
           client.get("/api/v1/contracts", headers={"Authorization": ""}).status_code)
    record("cross-owner contract read", 403,
           client.get(f"/api/v1/contracts/{contract.id}").status_code)
    response = client.get(f"/api/v1/contracts/versions/{version.id}/clauses")
    record("cross-owner clause read", 403, response.status_code)
    record("unprivileged prompt edit", 403,
           client.put(f"/api/v1/prompts/{prompt.id}/content", json={"content": "altered by ordinary user"}).status_code)
    record("cross-user QA trace read", 403,
           client.get("/api/v1/agents/trace?task_type=qa&ref_id=0").status_code)
    # eligibility/amendment scenario: alice owns this contract (scope) and may start approvals
    # (workflow:start), so the remaining gates under test are the report checks and the amend loop
    own = helpers.make_contract(db, owner_id=alice.id, status="review_passed")
    own_version = helpers.make_version(db, own.id, parse_status="parsed")
    own_task = helpers.make_review_task(db, own.id, version_id=own_version.id, status="completed")
    own_clause = helpers.make_clause(db, own_version.id, own.id, content="第一条 付款：甲方支付价款。")
    helpers.make_report(db, own_task.id, own.id, score=100,
                        report_json={"summary": {"failedClauses": 1}})
    db.commit()
    record("incomplete report approval", 400,
           client.post("/api/v1/workflow/instances", json={"contractId": own.id}).status_code)
    record("unprivileged workflow edit", 403,
           client.post("/api/v1/workflow/definitions", json={"name": "ordinary edit", "contractTypes": ["purchase"],
                       "nodes": [{"id": "cc", "type": "cc", "name": "only cc"}]}).status_code)
    with patch("app.services.contract_service.index_clauses_to_es", lambda *a, **k: None):
        response = client.put(f"/api/v1/findings/{finding.id}/status", json={"status": "accepted"})
    db.expire_all()
    record("cross-owner finding mutation", 403, response.status_code)
    # 采纳建议（本人合同）：必须换版或作废旧审查结论
    own_finding = helpers.make_finding(db, own_task.id, clause_id=own_clause.id, suggested_text="changed clause")
    db.commit()
    with patch("app.services.contract_service.index_clauses_to_es", lambda *a, **k: None):
        client.put(f"/api/v1/findings/{own_finding.id}/status", json={"status": "accepted"})
    db.expire_all()
    record("amendment invalidates report/version", True,
           db.get(type(own), own.id).current_version != 1 or
           db.get(type(own_task), own_task.id).status != "completed")

    graph = {"nodes": [{"id": "cond", "type": "condition", "rules": [{"when": "default", "goto": "approval"}]},
                       {"id": "approval", "type": "approval", "name": "mandatory", "assignees": ["role:a"]}],
             "edges": [{"from": "start", "to": "cond"}, {"from": "approval", "to": "end"}]}
    branch_def = helpers.make_definition(db, name="branch", graph=graph)
    instance = helpers.make_instance(db, branch_def.id, contract.id)
    db.commit()
    workflow_engine._advance_from(db, instance, branch_def, "start")
    record("condition must stop at selected approval", "approval", instance.current_node_id)
    # Real act(): two users of role a can satisfy a node that also requires role b.
    sign_graph = helpers.approval_graph(nodes=[{"id": "n1", "type": "approval", "signMode": "and",
                                               "assignees": ["role:a", "role:b"]}])
    sign_def = helpers.make_definition(db, name="sign", graph=sign_graph)
    sign_inst = helpers.make_instance(db, sign_def.id, contract.id, current_node_id="n1")
    role_a = helpers.make_role(db, code="a")
    a1 = helpers.make_user(db, username="aone", real_name="A one", roles=[role_a])
    a2 = helpers.make_user(db, username="atwo", real_name="A two", roles=[role_a])
    db.commit()
    # After the fix the second role-a user is rejected (slot for role:b stays open) instead of
    # wrongly completing the node — treat the rejection as the correct outcome.
    for operator in (a1, a2):
        try:
            workflow_engine.act(db, sign_inst, action="approve", operator=operator, comment="")
        except Exception:
            pass
    record("and sign requires role b approval", "running", sign_inst.status)

    path = infra._TMP_ROOT / "table.docx"
    document = Document()
    document.add_paragraph("第一条 正文")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "critical table term"
    document.save(path)
    pages, _ = parse_service._parse_docx(path)
    record("DOCX table extracted", True, "critical table term" in "".join(pages))
    try:
        review_service.LlmClauseResult.model_validate({"risk_level": "HIGH"})
        rejected_invalid_level = False
    except Exception:
        rejected_invalid_level = True
    record("invalid model risk level rejected", True, rejected_invalid_level)
    penalty_clause = ContractClause(id=999, clause_no="test", clause_type="payment",
                                    content="预付款50%，逾期违约金为合同金额5%。发生争议提交仲裁。")
    findings = review_service.run_rule_engine(db, contract, [penalty_clause])
    record("payment percentage must not become penalty percentage", False,
           any("违约金比例 50%" in finding["issue"] for finding in findings))
    client.close()
    app.dependency_overrides.clear()
    db.rollback()
    db.close()
    output = Path(__file__).resolve().parent.parent / "audit-results-2026-10-21.json"
    output.write_text(json.dumps(results, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(results, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    faulthandler.dump_traceback_later(25, repeat=False)
    with patch.object(socket.socket, "connect", no_network), patch.object(socket.socket, "connect_ex", no_network_ex):
        main()
