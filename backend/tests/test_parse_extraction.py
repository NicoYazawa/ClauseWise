"""Batch 4 regression: document extraction and rule-engine counterexamples (audit counterexamples must all be handled correctly).

- DOCX tables must enter the parse result (amounts/payment/liability often live in tables)
- PDF keeps original page numbers and reading order (mixed scanned pages must not compress page numbers or be appended at the end)
- strict model risk-level validation; percentages bound to specific obligations; rule citations must not forge retrieval similarity
"""
from __future__ import annotations

from pathlib import Path

import pytest

from app.models.contract import Contract, ContractClause
from app.services import parse_service, review_service
from tests import helpers


# ===================== DOCX tables =====================
class TestDocxExtraction:
    def test_docx_tables_extracted_in_document_order(self, tmp_path):
        """Audit counterexample: a DOCX with body + table clauses must have its table content parsed."""
        from docx import Document

        path = tmp_path / "table.docx"
        document = Document()
        document.add_paragraph("第一条 正文")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "critical table term"
        document.add_paragraph("第二条 结尾")
        document.save(path)

        pages, channel = parse_service._parse_docx(path)
        text = "\n".join(pages)
        assert channel == "text"
        assert "critical table term" in text
        assert text.index("第一条 正文") < text.index("critical table term") < text.index("第二条 结尾")


# ===================== PDF page numbers and order =====================
class TestPdfPageOrder:
    def test_pdf_keeps_original_page_numbers_and_reading_order(self, db, monkeypatch, tmp_path):
        import pymupdf as fitz

        user = helpers.make_user(db)
        contract = helpers.make_contract(db, owner_id=user.id, status="parsing", current_version=1)
        version = helpers.make_version(db, contract.id, version_no=1, file_path="mixed.pdf",
                                       parse_status="pending")
        db.commit()

        doc = fitz.open()
        p1 = doc.new_page()
        for row in range(3):  # multi-line text (default line width truncates; cumulative >50 chars counts as a text page)
            p1.insert_text((72, 72 + row * 16), "第一条 付款：甲方应于收到发票后三十日内支付全部价款。",
                           fontname="china-s")
        doc.new_page()  # blank scanned page (page 2)
        p3 = doc.new_page()
        for row in range(3):
            p3.insert_text((72, 72 + row * 16), "第二条 交付：甲方应于合同生效后三十日内交付标的物。",
                           fontname="china-s")
        pdf_path = Path(parse_service.settings.storage_path) / "mixed.pdf"
        pdf_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(pdf_path))
        doc.close()

        def fake_vision(path, page_no, **kw):
            return [{"clause_no": "扫描条款", "clause_type": "其他约定",
                     "content": f"OCR 第{page_no}页关键条款", "page_no": page_no}]

        monkeypatch.setattr(parse_service, "_vision_page", fake_vision)
        monkeypatch.setattr(parse_service, "index_clauses_to_es", lambda *a, **k: None)

        parse_service.parse_version(db, version.id)
        db.expire_all()
        clauses = (
            db.query(ContractClause)
            .filter(ContractClause.version_id == version.id)
            .order_by(ContractClause.id)
            .all()
        )
        assert clauses, "解析结果不得为空"
        # original page numbers preserved: text pages 1/3 and scanned page 2 stay in place
        assert [c.page_no for c in clauses] == sorted(c.page_no for c in clauses)
        assert {c.page_no for c in clauses} == {1, 2, 3}
        # reading order: OCR page content must not be appended to the document end
        contents = [c.content for c in clauses]
        ocr_idx = next(i for i, c in enumerate(contents) if "OCR 第2页" in c)
        page3_idx = next(i for i, c in enumerate(contents) if "第二条" in c)
        assert ocr_idx < page3_idx
        assert any("第一条" in c for c in contents)


# ===================== Rule-engine counterexamples =====================
class TestRuleEngineAnchors:
    def _findings(self, db, content: str):
        user = helpers.make_user(db)
        contract = helpers.make_contract(db, owner_id=user.id, contract_type="sales")
        clause = helpers.make_clause(db, 1, contract.id, clause_no="第一条", clause_type="付款",
                                     content=content)
        return review_service.run_rule_engine(db, contract, [clause])

    def test_unrelated_percentage_not_bound_to_penalty(self, db):
        """Audit counterexample: a 50% advance payment + 5% penalty must not yield "penalty ratio 50%"."""
        findings = self._findings(
            db, "预付款50%，逾期违约金为合同金额5%。发生争议提交仲裁。")
        assert not [f for f in findings if "违约金比例 50%" in f["issue"]]
        assert not [f for f in findings if "违约金比例" in f["issue"]]

    def test_penalty_percentage_still_flagged(self, db):
        findings = self._findings(db, "违约金为合同总价的 50%。")
        assert any("违约金比例 50%" in f["issue"] for f in findings)

    def test_rule_findings_carry_rule_citations_not_fake_similarity(self, db):
        findings = self._findings(db, "甲方有权随时解除合同且不承担任何责任。")
        assert findings
        for f in findings:
            for hit in f.get("rag_hits") or []:
                assert "相似度" not in hit, "规则依据不得伪造检索相似度"
                assert "规则引用" in hit

    def test_invalid_model_risk_output_is_rejected(self):
        """Audit counterexample: an illegal risk level must fail validation, not be treated as no-risk."""
        from pydantic import ValidationError

        with pytest.raises(ValidationError):
            review_service.LlmClauseResult.model_validate({"risk_level": "HIGH"})
        with pytest.raises(ValidationError):
            review_service.LlmClauseResult.model_validate({"risk_level": "high", "risk_points": []})
