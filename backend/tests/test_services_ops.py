"""Unit tests for the services layer (part 2): parsing / knowledge base / Q&A / diff summarization / report export."""
from __future__ import annotations

import io
from datetime import date
from types import SimpleNamespace

import pytest

from app.ai import llm_client
from app.core.exceptions import AppError
from app.services import (
    diff_llm_service,
    knowledge_service,
    parse_service,
    qa_service,
    report_export,
)
from tests import helpers


def _write_file(storage_root, rel: str, data: bytes) -> str:
    target = storage_root / rel
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(data)
    return rel


# ===================== parse_service =====================
class TestParseService:
    def test_parse_docx(self, tmp_path):
        from docx import Document as DocxDocument

        doc = DocxDocument()
        doc.add_paragraph("第一条 付款方式")
        doc.add_paragraph("甲方应在收到发票后三十日内支付全部价款人民币一百万元。")
        p = tmp_path / "t.docx"
        doc.save(str(p))
        pages, channel = parse_service._parse_docx(p)
        assert channel == "text"
        assert len(pages) == 1
        assert "第一条" in pages[0]

    def test_parse_docx_pagination(self, tmp_path):
        from docx import Document as DocxDocument

        doc = DocxDocument()
        for i in range(4):
            doc.add_paragraph(f"第{i}条 " + "内容" * 200)
        p = tmp_path / "t.docx"
        doc.save(str(p))
        pages, _ = parse_service._parse_docx(p)
        assert len(pages) >= 2  # approximate pagination

    def test_parse_pdf_text_and_scanned(self, tmp_path):
        import pymupdf as fitz

        doc = fitz.open()
        # PyMuPDF's built-in font does not render Chinese; use ASCII text for extractable content
        doc.new_page().insert_text((72, 72), "Payment terms of article 1 " + "content " * 20)
        doc.new_page()  # blank page -> scanned page
        p = tmp_path / "t.pdf"
        doc.save(str(p))
        doc.close()
        pages, scan_pages, channel = parse_service._parse_pdf(p)
        assert len(pages) == 2
        assert scan_pages == [2]
        assert channel == "mixed"
        assert "Payment terms" in pages[0]

    def test_parse_pdf_all_scanned(self, tmp_path):
        import pymupdf as fitz

        doc = fitz.open()
        doc.new_page()
        p = tmp_path / "t.pdf"
        doc.save(str(p))
        doc.close()
        _, scan_pages, channel = parse_service._parse_pdf(p)
        assert scan_pages == [1]
        assert channel == "vision"

    def test_render_pdf_page_png(self, tmp_path):
        import pymupdf as fitz

        doc = fitz.open()
        doc.new_page()
        p = tmp_path / "t.pdf"
        doc.save(str(p))
        doc.close()
        assert parse_service._render_pdf_page_png(p, 1)[:4] == b"\x89PNG"

    def test_vision_page_image(self, tmp_path, monkeypatch):
        import pymupdf as fitz

        pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 20, 20))
        p = tmp_path / "scan.png"
        p.write_bytes(pix.tobytes("png"))
        monkeypatch.setattr(llm_client, "vision_extract_clauses",
                            lambda b64, mime="image/png", prompt=None: ({"clauses": [{"clause_no": "第一条"}]}, 5))
        assert parse_service._vision_page(p, 1, is_image=True) == [{"clause_no": "第一条", "page_no": 1}]

    def test_vision_page_pdf_render(self, tmp_path, monkeypatch):
        import pymupdf as fitz

        doc = fitz.open()
        doc.new_page()
        p = tmp_path / "t.pdf"
        doc.save(str(p))
        doc.close()
        monkeypatch.setattr(llm_client, "vision_extract_clauses",
                            lambda b64, mime="image/png", prompt=None: ({"clauses": []}, 5))
        assert parse_service._vision_page(p, 1) == []

    def _make_docx(self) -> bytes:
        from docx import Document as DocxDocument

        doc = DocxDocument()
        doc.add_paragraph("第一条 付款方式：甲方应在收到发票后三十日内支付全部价款。")
        doc.add_paragraph("第二条 违约责任：违约方应支付合同总价百分之十的违约金。")
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    def test_parse_version_docx(self, db, storage_root, monkeypatch):
        from app.models.contract import ContractClause

        user = helpers.make_user(db)
        contract = helpers.make_contract(db, owner_id=user.id, status="parsing")
        version = helpers.make_version(db, contract.id, parse_status="pending")
        version.file_path = _write_file(storage_root, "contracts/1/origin/v1_a.docx", self._make_docx())
        db.commit()
        indexed = []
        monkeypatch.setattr(parse_service, "index_clauses_to_es", lambda db, ver: indexed.append(ver.id))
        parse_service.parse_version(db, version.id)
        db.expire_all()
        ver = db.get(type(version), version.id)
        assert ver.parse_status == "parsed"
        assert ver.parse_channel == "text"
        assert db.query(ContractClause).filter_by(version_id=version.id).count() == 2
        assert indexed == [version.id]
        assert db.get(type(contract), contract.id).status == "parsed"

    def test_parse_version_pdf_mixed(self, db, storage_root, monkeypatch):
        import pymupdf as fitz

        from app.models.contract import ContractClause

        user = helpers.make_user(db)
        contract = helpers.make_contract(db, owner_id=user.id, status="parsing")
        version = helpers.make_version(db, contract.id, parse_status="pending")
        doc = fitz.open()
        doc.new_page().insert_text((72, 72), "Payment terms of article 1 " + "content text " * 10)
        doc.new_page()  # scanned page
        buf = io.BytesIO()
        doc.save(buf)
        doc.close()
        version.file_path = _write_file(storage_root, "contracts/1/origin/v1_b.pdf", buf.getvalue())
        db.commit()
        monkeypatch.setattr(parse_service, "index_clauses_to_es", lambda db, ver: None)
        monkeypatch.setattr(llm_client, "vision_extract_clauses",
                            lambda b64, mime="image/png", prompt=None: (
                                {"clauses": [{"clause_no": "扫描第一条", "content": "扫描出来的条款内容。"}]}, 5))
        parse_service.parse_version(db, version.id)
        db.expire_all()
        assert db.get(type(version), version.id).parse_channel == "mixed"
        nos = {c.clause_no for c in db.query(ContractClause).filter_by(version_id=version.id).all()}
        assert nos == {"未编号", "扫描第一条"}  # ASCII text page has no Chinese clause number -> unnumbered; scanned page uses the vision channel

    def test_parse_version_image(self, db, storage_root, monkeypatch):
        import pymupdf as fitz

        user = helpers.make_user(db)
        contract = helpers.make_contract(db, owner_id=user.id, status="parsing")
        version = helpers.make_version(db, contract.id, parse_status="pending")
        pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 20, 20))
        version.file_path = _write_file(storage_root, "contracts/1/origin/v1_c.png", pix.tobytes("png"))
        db.commit()
        monkeypatch.setattr(parse_service, "index_clauses_to_es", lambda db, ver: None)
        monkeypatch.setattr(llm_client, "vision_extract_clauses",
                            lambda b64, mime="image/png", prompt=None: (
                                {"clauses": [{"clause_no": "第一条", "content": "图片条款内容。"}]}, 5))
        parse_service.parse_version(db, version.id)
        db.expire_all()
        assert db.get(type(version), version.id).parse_channel == "vision"

    def test_parse_version_missing(self, db):
        with pytest.raises(parse_service.ParseError):
            parse_service.parse_version(db, 99999)

    def test_parse_version_already_claimed(self, db):
        user = helpers.make_user(db)
        contract = helpers.make_contract(db, owner_id=user.id)
        version = helpers.make_version(db, contract.id, parse_status="parsed")
        db.commit()
        parse_service.parse_version(db, version.id)  # returns directly
        db.expire_all()
        assert db.get(type(version), version.id).parse_status == "parsed"

    def test_parse_version_no_file_path(self, db):
        user = helpers.make_user(db)
        contract = helpers.make_contract(db, owner_id=user.id, status="parsing")
        version = helpers.make_version(db, contract.id, parse_status="pending", file_path=None)
        db.commit()
        with pytest.raises(parse_service.ParseError) as ei:
            parse_service.parse_version(db, version.id)
        assert "文件路径缺失" in str(ei.value)
        db.expire_all()
        assert db.get(type(version), version.id).parse_status == "failed"
        assert db.get(type(contract), contract.id).status == "parse_failed"

    def test_parse_version_file_missing(self, db):
        user = helpers.make_user(db)
        contract = helpers.make_contract(db, owner_id=user.id, status="parsing")
        version = helpers.make_version(db, contract.id, parse_status="pending",
                                       file_path="contracts/1/origin/v1_gone.pdf")
        db.commit()
        with pytest.raises(parse_service.ParseError) as ei:
            parse_service.parse_version(db, version.id)
        assert "原始文件不存在" in str(ei.value)

    def test_parse_version_unsupported_ext(self, db, storage_root):
        user = helpers.make_user(db)
        contract = helpers.make_contract(db, owner_id=user.id, status="parsing")
        version = helpers.make_version(db, contract.id, parse_status="pending",
                                       file_path=_write_file(storage_root,
                                                             "contracts/1/origin/v1_x.txt", b"hello"))
        db.commit()
        with pytest.raises(parse_service.ParseError) as ei:
            parse_service.parse_version(db, version.id)
        assert "不支持的文件格式" in str(ei.value)

    def test_parse_version_no_clauses(self, db, storage_root):
        from docx import Document as DocxDocument

        user = helpers.make_user(db)
        contract = helpers.make_contract(db, owner_id=user.id, status="parsing")
        version = helpers.make_version(db, contract.id, parse_status="pending")
        doc = DocxDocument()
        doc.add_paragraph("   ")  # empty paragraph -> no clause
        buf = io.BytesIO()
        doc.save(buf)
        version.file_path = _write_file(storage_root, "contracts/1/origin/v1_e.docx", buf.getvalue())
        db.commit()
        with pytest.raises(parse_service.ParseError) as ei:
            parse_service.parse_version(db, version.id)
        assert "未解析出任何有效条款" in str(ei.value)

    def test_parse_version_retry_clears_old_clauses(self, db, storage_root, monkeypatch):
        from app.models.contract import ContractClause

        user = helpers.make_user(db)
        contract = helpers.make_contract(db, owner_id=user.id, status="parsing")
        version = helpers.make_version(db, contract.id, parse_status="pending")
        helpers.make_clause(db, version.id, contract.id, clause_no="旧条款", content="旧内容。")
        version.file_path = _write_file(storage_root, "contracts/1/origin/v1_r.docx", self._make_docx())
        db.commit()
        monkeypatch.setattr(parse_service, "index_clauses_to_es", lambda db, ver: None)
        parse_service.parse_version(db, version.id)
        db.expire_all()
        clauses = db.query(ContractClause).filter_by(version_id=version.id).all()
        assert [c.clause_no for c in clauses] == ["第一条", "第二条"]


# ===================== knowledge_service =====================
class TestKnowledgeService:
    def test_parse_date(self):
        assert knowledge_service._parse_date("2026-01-01") == date(2026, 1, 1)
        assert knowledge_service._parse_date("bad") is None
        assert knowledge_service._parse_date(None) is None
        assert knowledge_service._parse_date("") is None

    def test_index_law(self, db, monkeypatch):
        es = helpers.FakeES()
        monkeypatch.setattr(knowledge_service, "get_es_client", lambda: es)
        monkeypatch.setattr(knowledge_service, "embed_texts", lambda texts: [[0.1] * 4])
        law = helpers.make_law(db, title="民法典", content="合同编")
        db.commit()
        assert knowledge_service.index_law(db, law) is True
        db.expire_all()
        assert law.vectorized is True
        assert law.es_doc_id == f"law_{law.id}"
        assert es.indexed[0]["document"]["content_vector"] == [0.1] * 4

    def test_index_law_no_es(self, db, monkeypatch):
        monkeypatch.setattr(knowledge_service, "get_es_client", lambda: None)
        law = helpers.make_law(db)
        db.commit()
        assert knowledge_service.index_law(db, law) is False
        db.expire_all()
        assert law.vectorized is False

    def test_index_law_failure(self, db, monkeypatch):
        class _BoomES(helpers.FakeES):
            def index(self, **kw):
                raise RuntimeError("es down")

        monkeypatch.setattr(knowledge_service, "get_es_client", lambda: _BoomES())
        monkeypatch.setattr(knowledge_service, "embed_texts", lambda texts: [[0.1]])
        law = helpers.make_law(db)
        db.commit()
        assert knowledge_service.index_law(db, law) is False
        db.expire_all()
        assert law.vectorized is False

    def test_delete_law_from_es(self, db, monkeypatch):
        es = helpers.FakeES()
        monkeypatch.setattr(knowledge_service, "get_es_client", lambda: es)
        law = helpers.make_law(db, es_doc_id="law_1")
        db.commit()
        knowledge_service.delete_law_from_es(law)
        assert es.deleted == ["legal_knowledge/law_1"]

    def test_delete_law_from_es_no_client_or_id(self, db, monkeypatch):
        monkeypatch.setattr(knowledge_service, "get_es_client", lambda: None)
        knowledge_service.delete_law_from_es(helpers.make_law(db, es_doc_id="law_1"))
        monkeypatch.setattr(knowledge_service, "get_es_client", lambda: helpers.FakeES())
        knowledge_service.delete_law_from_es(helpers.make_law(db, es_doc_id=None))  # does not raise

    def test_index_template(self, db, monkeypatch):
        es = helpers.FakeES()
        monkeypatch.setattr(knowledge_service, "get_es_client", lambda: es)
        monkeypatch.setattr(knowledge_service, "embed_texts", lambda texts: [[0.2] * 4])
        tpl = helpers.make_template(db)
        db.commit()
        assert knowledge_service.index_template(db, tpl) is True
        assert tpl.es_doc_id == f"tpl_{tpl.id}_v{tpl.version}"

    def test_index_template_no_es(self, db, monkeypatch):
        monkeypatch.setattr(knowledge_service, "get_es_client", lambda: None)
        tpl = helpers.make_template(db)
        db.commit()
        assert knowledge_service.index_template(db, tpl) is False

    def test_index_template_failure(self, db, monkeypatch):
        class _BoomES(helpers.FakeES):
            def index(self, **kw):
                raise RuntimeError("es down")

        monkeypatch.setattr(knowledge_service, "get_es_client", lambda: _BoomES())
        monkeypatch.setattr(knowledge_service, "embed_texts", lambda texts: [[0.1]])
        tpl = helpers.make_template(db)
        db.commit()
        assert knowledge_service.index_template(db, tpl) is False
        db.expire_all()
        assert tpl.vectorized is False

    def test_delete_template_from_es(self, db, monkeypatch):
        es = helpers.FakeES()
        monkeypatch.setattr(knowledge_service, "get_es_client", lambda: es)
        tpl = helpers.make_template(db, es_doc_id="tpl_1_v1")
        db.commit()
        knowledge_service.delete_template_from_es(tpl)
        assert es.deleted == ["template_clause/tpl_1_v1"]

    def test_delete_template_from_es_noop(self, db, monkeypatch):
        monkeypatch.setattr(knowledge_service, "get_es_client", lambda: None)
        knowledge_service.delete_template_from_es(helpers.make_template(db))
        monkeypatch.setattr(knowledge_service, "get_es_client", lambda: helpers.FakeES())
        knowledge_service.delete_template_from_es(helpers.make_template(db, es_doc_id=None))

    def test_bulk_index_contract_clauses(self, monkeypatch):
        import elasticsearch.helpers as es_helpers

        calls = []
        monkeypatch.setattr(es_helpers, "bulk", lambda es, actions: calls.append(actions))
        monkeypatch.setattr(knowledge_service, "get_es_client", lambda: helpers.FakeES())
        knowledge_service.bulk_index_contract_clauses([{"clause_id": 1, "content": "c"}])
        assert calls[0][0]["_id"] == "clause_1"
        assert calls[0][0]["_index"] == "contract_index"

    def test_bulk_index_no_es_or_docs(self, monkeypatch):
        monkeypatch.setattr(knowledge_service, "get_es_client", lambda: None)
        knowledge_service.bulk_index_contract_clauses([{"clause_id": 1}])
        monkeypatch.setattr(knowledge_service, "get_es_client", lambda: helpers.FakeES())
        knowledge_service.bulk_index_contract_clauses([])

    def test_bulk_index_failure_logged(self, monkeypatch):
        import elasticsearch.helpers as es_helpers

        def _boom(es, actions):
            raise RuntimeError("bulk down")

        monkeypatch.setattr(es_helpers, "bulk", _boom)
        monkeypatch.setattr(knowledge_service, "get_es_client", lambda: helpers.FakeES())
        knowledge_service.bulk_index_contract_clauses([{"clause_id": 1}])  # does not raise

    def test_parse_law_rows(self):
        rows = [
            {"docType": "case", "title": " 案例一 ", "content": "内容", "sourceRef": "src",
             "effectiveDate": "2026-01-01", "expiryDate": "2027-01-01", "status": "active"},
            {"doc_type": "law", "title": "法条", "content": "正文", "status": "expired"},  # snake_case
            {"docType": "bad", "title": "t", "content": "c"},  # illegal type -> law
            {"docType": "law", "title": "t", "content": "c", "status": "bad"},  # illegal status -> active
            {"docType": "law", "title": "", "content": "c"},  # no title -> filtered
            {"docType": "law", "title": "t", "content": ""},  # no content -> filtered
            {"docType": "law", "title": "t", "content": "c", "effectiveDate": "bad-date"},  # bad date -> None
        ]
        laws = knowledge_service.parse_law_rows(rows)
        assert len(laws) == 5
        assert laws[0].doc_type == "case" and laws[0].title == "案例一"
        assert laws[2].doc_type == "law"
        assert laws[3].status == "active"
        assert laws[4].effective_date is None

    def test_read_law_excel(self, tmp_path):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.append(["类型", "标题", "内容", "来源", "生效日期", "失效日期", "状态"])
        ws.append(["law", "民法典", "合同编内容", "官网", "2026-01-01", None, "active"])
        ws.append([None, None, None, None, None, None, None])  # empty row skipped
        ws.append(["case", "指导案例", "案例内容", "官网", None, None, "active"])
        p = tmp_path / "laws.xlsx"
        wb.save(str(p))
        rows = knowledge_service.read_law_excel(str(p))
        assert len(rows) == 2
        assert rows[0]["docType"] == "law"
        assert rows[0]["title"] == "民法典"

    def test_read_law_excel_empty(self, tmp_path):
        from openpyxl import Workbook

        wb = Workbook()
        p = tmp_path / "empty.xlsx"
        wb.save(str(p))
        assert knowledge_service.read_law_excel(str(p)) == []


# ===================== qa_service =====================
class TestQaService:
    def test_ask_empty(self, db):
        with pytest.raises(AppError) as ei:
            qa_service.ask(db, "   ")
        assert ei.value.detail["code"] == "QA_EMPTY_QUESTION"

    def test_ask(self, db, monkeypatch):
        monkeypatch.setattr(llm_client, "ensure_llm_configured", lambda: None)
        monkeypatch.setattr("app.agents.qa_graph.run_qa_graph",
                            lambda q, **kwargs: {"answer": "答案",
                                       "citations": [{"kind": "law", "id": 1, "title": "t",
                                                      "snippet": "s", "score": 0.9}],
                                       "relevant_hits": [{"kind": "law", "ref_id": 1}], "no_hit": False})
        law = helpers.make_law(db)
        db.commit()
        out = qa_service.ask(db, "违约金")
        assert out["answer"] == "答案"
        db.expire_all()
        assert law.hit_count == 1  # hit count +1

    def test_ask_no_hit_no_bump(self, db, monkeypatch):
        monkeypatch.setattr(llm_client, "ensure_llm_configured", lambda: None)
        monkeypatch.setattr("app.agents.qa_graph.run_qa_graph",
                            lambda q, **kwargs: {"answer": "引导", "citations": [], "relevant_hits": [],
                                       "no_hit": True})
        assert qa_service.ask(db, "问题")["citations"] == []

    def test_ask_stream(self, db, monkeypatch):
        events = [
            {"type": "agent_step", "agent": "qa_router", "role": "意图路由", "summary": "s", "detail": {}},
            {"type": "citations", "citations": []},
            {"type": "delta", "content": "答"},
            {"type": "_hits", "hits": [{"kind": "law", "ref_id": 1}]},
            {"type": "done"},
        ]
        monkeypatch.setattr("app.agents.qa_graph.stream_qa_graph", lambda q, **kwargs: iter(events))
        law = helpers.make_law(db)
        db.commit()
        out = list(qa_service.ask_stream(db, "问题"))
        assert [e["type"] for e in out] == ["agent_step", "citations", "delta", "done"]
        db.expire_all()
        assert law.hit_count == 1

    def test_bump_hit_counts_failure_swallowed(self):
        class _FakeSession:
            def __init__(self):
                self.rolled = False

            def get(self, model, pk):
                return SimpleNamespace(hit_count=1)

            def commit(self):
                raise RuntimeError("db down")

            def rollback(self):
                self.rolled = True

        fake = _FakeSession()
        qa_service._bump_hit_counts(fake, [{"kind": "law", "ref_id": 1}])
        assert fake.rolled is True

    def test_bump_hit_counts_ignores_non_law(self, db):
        law = helpers.make_law(db)
        db.commit()
        qa_service._bump_hit_counts(db, [{"kind": "template", "ref_id": law.id},
                                         {"kind": "law", "ref_id": None}])
        db.expire_all()
        assert law.hit_count == 0


# ===================== diff_llm_service =====================
class TestDiffLlmService:
    def test_summarize_diff_success(self, monkeypatch):
        monkeypatch.setattr(llm_client, "chat_json", lambda s, u: ({"kind": "drift", "summary": "语义漂移"}, 5))
        kind, summary = diff_llm_service.summarize_diff("旧", "新")
        assert kind == "drift" and summary == "语义漂移"

    def test_summarize_diff_invalid_kind(self, monkeypatch):
        monkeypatch.setattr(llm_client, "chat_json", lambda s, u: ({"kind": "weird"}, 5))
        assert diff_llm_service.summarize_diff("旧", "新")[0] == "modified"

    def test_summarize_diff_empty_summary_fallback(self, monkeypatch):
        monkeypatch.setattr(llm_client, "chat_json", lambda s, u: ({"kind": "modified", "summary": "  "}, 5))
        kind, summary = diff_llm_service.summarize_diff("旧条款", "新条款")
        assert kind == "modified"
        assert summary  # difflib fallback

    def test_summarize_diff_truncates(self, monkeypatch):
        monkeypatch.setattr(llm_client, "chat_json", lambda s, u: ({"kind": "modified", "summary": "总" * 500}, 5))
        assert len(diff_llm_service.summarize_diff("旧", "新")[1]) == 300

    def test_summarize_diff_llm_failure(self, monkeypatch):
        def _boom(s, u):
            raise llm_client.LLMError("down")

        monkeypatch.setattr(llm_client, "chat_json", _boom)
        kind, summary = diff_llm_service.summarize_diff("旧条款", "新条款")
        assert kind == "modified"
        assert "改为" in summary or "一致" in summary


# ===================== report_export =====================
class TestReportExport:
    def _objects(self):
        from app.models.review import ReviewReport, ReviewTask, RiskFinding

        task = ReviewTask(id=1, contract_id=1, task_type="full", status="completed",
                          model_version="deepseek-chat", prompt_version="v1.0",
                          duration_ms=1234, tokens_used=100)
        report = ReviewReport(id=1, task_id=1, contract_id=1, score=90, high_count=1,
                              medium_count=1, low_count=1, overall_opinion="整体意见")
        findings = [
            RiskFinding(id=1, task_id=1, clause_no="第一条", risk_level="high", risk_type="权责不对等",
                        issue="问题", original_text="原文", legal_basis="依据", suggestion="建议",
                        suggested_text="改写", source="rule"),
            RiskFinding(id=2, task_id=1, clause_no="第二条", risk_level="medium", risk_type="违约",
                        issue="问题2", original_text="", legal_basis="", suggestion="", source="llm"),
        ]
        return report, task, findings

    def test_find_cjk_font(self):
        font = report_export._find_cjk_font()
        assert font is None or font.endswith((".ttc", ".ttf"))

    def test_export_word(self):
        report, task, findings = self._objects()
        data = report_export.export_word(report, task, findings)
        assert data[:2] == b"PK"  # docx = zip
        assert len(data) > 1000

    def test_export_word_no_suggested_text(self):
        from app.models.review import ReviewReport, ReviewTask, RiskFinding

        task = ReviewTask(id=1, contract_id=1, task_type="full", model_version="m", prompt_version="v")
        report = ReviewReport(id=1, task_id=1, contract_id=1, score=100, overall_opinion="通过")
        finding = RiskFinding(id=1, task_id=1, clause_no="第一条", risk_level="low", risk_type="x",
                              issue="i", original_text="o", legal_basis="l", suggestion="s",
                              suggested_text=None, source="llm")
        data = report_export.export_word(report, task, [finding])
        assert data[:2] == b"PK"

    def test_export_pdf(self):
        report, task, findings = self._objects()
        data = report_export.export_pdf(report, task, findings)
        assert data[:5] == b"%PDF-"
        assert len(data) > 1000

    def test_export_pdf_without_cjk_font(self, monkeypatch):
        monkeypatch.setattr(report_export, "_find_cjk_font", lambda: None)
        report, task, findings = self._objects()
        data = report_export.export_pdf(report, task, findings)
        assert data[:5] == b"%PDF-"  # falls back to Helvetica and still generates

    def test_level_labels(self):
        assert report_export.LEVEL_LABELS["high"] == "高风险"
        assert report_export.LEVEL_LABELS["none"] == "无风险"
